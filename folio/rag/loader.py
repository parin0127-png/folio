from pypdf import PdfReader
from docx import Document
from bs4 import BeautifulSoup
from pptx import Presentation
import pandas as pd
import json 


class DocumentLoader:
    def load(self, file_path):
        ext = file_path.split(".")[-1].lower()
        loader = {
            "txt": self._txt, "md": self._txt,
            "pdf": self._pdf, "docx": self._docx,
            "csv": self._csv, "json": self._json,
            "html": self._html, "xlsx": self._excel, 
            "pptx": self._pptx,
        }
        if ext not in loader:
            raise ValueError(f"Unsupported file type: {ext}")
        return loader[ext](file_path)

    def _txt(self, path):
        with open(path, "r", encoding = "utf-8")as f:
            return f.read()

    def _json(self, path):
        with open(path, "r", encoding = "utf-8")as f:
            return json.dumps(json.load(f), indent = 2)

    def _html(self, path):
        with open(path, "r" , encoding = "utf-8")as f:
            return BeautifulSoup(f.read(), "html.parser").get_text()

    def _pdf(self, path):
        reader = PdfReader(path)
        return "".join(page.extract_text() for page in reader.pages)

    def _docx(self, path):
        return "\n".join(p.text for p in Document(path).paragraphs)

    def _csv(self, path):
        return pd.read_csv(path).to_string(index = False)

    def _excel(self, path):
        return pd.read_excel(path).to_string(index = False)

    def _pptx(self, path):
        prs = Presentation(path)
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text += shape.text_frame.text + "\n"
        return text
    